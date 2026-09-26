# -*- coding: utf-8 -*-
"""
news_paper.py  —  시사주간뉴스 → 진짜 종이신문(A3 세로 · 8면) 발행 엔진
================================================================================
· PDF   : reportlab 캔버스에 직접 조판합니다. (실제 신문처럼 제호·단·괘선·사진)
· WORD  : python-docx 로 A3 세로 2단 문서를 만듭니다.
· 그림  : PIL 로 사진·만평·도해를 직접 그립니다. (인터넷이 막혀도 항상 나옵니다)

app.py 는 이 파일을 import 해서 쓰기만 합니다. 기존 기능은 전혀 건드리지 않습니다.
streamlit 을 import 하지 않으므로 단독 실행·시험이 가능합니다.
"""

import io
import os
import re
import math
import random
import datetime

from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps

from reportlab.lib.pagesizes import A3
from reportlab.lib.units import mm
from reportlab.lib import colors as rlcolors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.platypus import Paragraph, Frame, Table, TableStyle, Spacer
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

# ==============================================================================
# 0. 폰트
# ==============================================================================
_FONT_REG_PATH = None
_FONT_BOLD_PATH = None
_BG_PROVIDER = None          # app.py 의 get_background_bytes(seed, theme, size)

F_REG = "Helvetica"
F_BOLD = "Helvetica-Bold"


def configure(font_regular=None, font_bold=None, bg_provider=None):
    """app.py 에서 한 번 불러 주면 됩니다."""
    global _FONT_REG_PATH, _FONT_BOLD_PATH, _BG_PROVIDER, F_REG, F_BOLD
    if font_regular and os.path.exists(font_regular):
        _FONT_REG_PATH = font_regular
    if font_bold and os.path.exists(font_bold):
        _FONT_BOLD_PATH = font_bold
    if bg_provider:
        _BG_PROVIDER = bg_provider

    for name, path in (("NP-R", _FONT_REG_PATH), ("NP-B", _FONT_BOLD_PATH)):
        if not path:
            continue
        try:
            pdfmetrics.registerFont(TTFont(name, path))
        except Exception:
            pass
    reg_ok = "NP-R" in pdfmetrics.getRegisteredFontNames()
    bold_ok = "NP-B" in pdfmetrics.getRegisteredFontNames()
    F_REG = "NP-R" if reg_ok else "Helvetica"
    F_BOLD = "NP-B" if bold_ok else ("NP-R" if reg_ok else "Helvetica-Bold")
    return F_REG, F_BOLD


def _pil_font(size, bold=True):
    path = (_FONT_BOLD_PATH if bold else _FONT_REG_PATH) or _FONT_REG_PATH or _FONT_BOLD_PATH
    if path:
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            pass
    return ImageFont.load_default()


# ==============================================================================
# 1. 공용 유틸
# ==============================================================================
def esc(s):
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _clean(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s


def sw(text, font, size):
    try:
        return pdfmetrics.stringWidth(text, font, size)
    except Exception:
        return len(text) * size * 0.6


def wrap_pdf(text, font, size, maxw):
    """CJK 안전 줄바꿈 (한글은 글자 단위, 영문은 낱말 단위)"""
    text = _clean(text)
    if not text:
        return []
    lines, cur = [], ""
    token = ""
    for ch in text:
        if re.match(r"[A-Za-z0-9\-\.\,\'\"/]", ch):
            token += ch
            continue
        if token:
            if sw(cur + token, font, size) > maxw and cur:
                lines.append(cur.rstrip())
                cur = token
            else:
                cur += token
            token = ""
        if sw(cur + ch, font, size) > maxw and cur.strip():
            lines.append(cur.rstrip())
            cur = "" if ch == " " else ch
        else:
            cur += ch
    if token:
        if sw(cur + token, font, size) > maxw and cur:
            lines.append(cur.rstrip())
            cur = token
        else:
            cur += token
    if cur.strip():
        lines.append(cur.rstrip())
    return lines


def wrap_pil(draw, text, font, maxw):
    text = _clean(text)
    if not text:
        return []
    lines, cur = [], ""
    for ch in text:
        t = cur + ch
        if draw.textlength(t, font=font) > maxw and cur.strip():
            lines.append(cur)
            cur = "" if ch == " " else ch
        else:
            cur = t
    if cur.strip():
        lines.append(cur)
    return lines


# ==============================================================================
# 2. 그림 만들기 (사진 · 만평 · 도해)
# ==============================================================================
def _procedural_bg(seed, size):
    """
    인터넷 없이도 항상 나오는 '사진 같은' 그림.
    하늘 → 해/달 → 원경 실루엣 → 중경 → 근경 으로 겹쳐 그려
    흑백으로 바꿔도 깊이가 살아 있게 만듭니다.
    """
    rnd = random.Random(str(seed))
    W, H = size

    # ── 하늘 그라데이션
    skies = [
        [(18, 28, 56), (60, 88, 148), (168, 196, 232)],
        [(36, 24, 20), (132, 92, 58), (238, 206, 160)],
        [(16, 34, 30), (52, 96, 82), (186, 214, 198)],
        [(28, 26, 34), (92, 92, 104), (210, 212, 220)],
        [(40, 20, 36), (120, 62, 96), (240, 206, 214)],
    ]
    pal = rnd.choice(skies)
    sky = Image.new("RGB", (2, 256))
    px = sky.load()
    for y in range(256):
        t = y / 255 * (len(pal) - 1)
        i = int(t)
        k = t - i
        a, b = pal[i], pal[min(i + 1, len(pal) - 1)]
        col = (int(a[0] + (b[0] - a[0]) * k), int(a[1] + (b[1] - a[1]) * k),
               int(a[2] + (b[2] - a[2]) * k))
        px[0, y] = px[1, y] = col
    img = sky.resize((W, H), Image.BICUBIC)
    d = ImageDraw.Draw(img, "RGBA")

    # ── 해 / 달
    sx, sy = rnd.randint(int(W * 0.12), int(W * 0.88)), rnd.randint(int(H * 0.12), int(H * 0.4))
    sr = rnd.randint(int(min(W, H) * 0.05), int(min(W, H) * 0.1))
    glow = Image.new("L", (W, H), 0)
    gd = ImageDraw.Draw(glow)
    gd.ellipse([sx - sr * 3.2, sy - sr * 3.2, sx + sr * 3.2, sy + sr * 3.2], fill=70)
    gd.ellipse([sx - sr * 1.5, sy - sr * 1.5, sx + sr * 1.5, sy + sr * 1.5], fill=150)
    gd.ellipse([sx - sr, sy - sr, sx + sr, sy + sr], fill=238)
    glow = glow.filter(ImageFilter.GaussianBlur(sr * 0.75))
    img = Image.composite(Image.new("RGB", (W, H), (255, 250, 236)), img, glow)
    d = ImageDraw.Draw(img, "RGBA")

    # ── 구름 띠
    for _ in range(rnd.randint(3, 7)):
        cy = rnd.randint(int(H * 0.06), int(H * 0.52))
        cw = rnd.randint(W // 5, W // 2)
        cx = rnd.randint(-cw // 3, W)
        ch = rnd.randint(H // 40, H // 16)
        d.ellipse([cx, cy, cx + cw, cy + ch],
                  fill=(255, 255, 255, rnd.randint(22, 52)))

    horizon = int(H * rnd.uniform(0.54, 0.7))
    kind = rnd.choice(["city", "hills", "forest", "sea"])

    def layer(top_y, colr, jag, step, alpha=255):
        pts = [(0, H)]
        x = 0
        y = top_y
        while x <= W:
            pts.append((x, y))
            x += step
            y = max(top_y - jag, min(top_y + jag, y + rnd.randint(-jag, jag)))
        pts.append((W, H))
        d.polygon(pts, fill=colr + (alpha,))

    if kind == "city":
        for li, (base, dark) in enumerate([(horizon - int(H * 0.10), 92),
                                           (horizon - int(H * 0.04), 58),
                                           (horizon + int(H * 0.02), 28)]):
            x = -20
            while x < W + 20:
                bw = rnd.randint(int(W * 0.03), int(W * 0.09))
                bh = rnd.randint(int(H * 0.04), int(H * 0.26)) * (1 - li * 0.22)
                top = base - bh
                d.rectangle([x, top, x + bw, H], fill=(dark, dark + 4, dark + 10))
                if li >= 1:
                    for wy in range(int(top) + 8, H, 16):
                        for wx in range(int(x) + 6, int(x + bw) - 6, 12):
                            if rnd.random() < 0.28:
                                d.rectangle([wx, wy, wx + 4, wy + 6],
                                            fill=(228, 216, 176, 190))
                x += bw + rnd.randint(3, 12)
    elif kind == "hills":
        layer(horizon - int(H * 0.12), (104, 112, 124), int(H * 0.05), W // 12)
        layer(horizon - int(H * 0.04), (64, 72, 84), int(H * 0.04), W // 14)
        layer(horizon + int(H * 0.04), (30, 34, 42), int(H * 0.03), W // 16)
    elif kind == "forest":
        layer(horizon - int(H * 0.02), (72, 80, 74), int(H * 0.02), W // 18)
        for li, (base, dark, sc) in enumerate([(horizon + int(H * 0.02), 46, 0.9),
                                               (H, 22, 1.25)]):
            x = -20
            while x < W + 20:
                tw = rnd.randint(int(W * 0.02), int(W * 0.05))
                th = int(rnd.randint(int(H * 0.08), int(H * 0.22)) * sc)
                d.polygon([(x, base), (x + tw / 2, base - th), (x + tw, base)],
                          fill=(dark, dark + 8, dark + 2))
                x += tw * rnd.uniform(0.45, 0.85)
    else:  # sea
        d.rectangle([0, horizon, W, H], fill=(38, 52, 74))
        for i in range(140):
            yy = horizon + (H - horizon) * (i / 140) ** 1.5
            ww = rnd.randint(W // 20, W // 5)
            xx = rnd.randint(0, W)
            d.line([(xx, yy), (xx + ww, yy)],
                   fill=(220, 226, 238, rnd.randint(20, 70)),
                   width=max(1, int(1 + (yy - horizon) / (H - horizon) * 4)))
        layer(horizon - int(H * 0.05), (48, 54, 64), int(H * 0.02), W // 14, 235)

    # ── 근경 사람 실루엣 (있을 때 사진처럼 보인다)
    if rnd.random() < 0.72:
        for _ in range(rnd.randint(1, 4)):
            fx = rnd.randint(int(W * 0.06), int(W * 0.94))
            fh = rnd.randint(int(H * 0.16), int(H * 0.3))
            fw = fh * 0.3
            fy = H - rnd.randint(0, int(H * 0.06))
            hr = fw * 0.34
            d.ellipse([fx - hr, fy - fh, fx + hr, fy - fh + hr * 2], fill=(14, 15, 18))
            d.polygon([(fx - fw * 0.5, fy), (fx - fw * 0.36, fy - fh + hr * 1.9),
                       (fx + fw * 0.36, fy - fh + hr * 1.9), (fx + fw * 0.5, fy)],
                      fill=(14, 15, 18))

    img = img.filter(ImageFilter.GaussianBlur(max(0.6, W / 900)))
    return img


def make_photo(seed, size=(1200, 800), theme="도시 · 일상", mono=True):
    """
    신문 사진 한 장. (배경 엔진이 있으면 그것을 쓰고, 없으면 직접 그립니다)
    신문 느낌을 위해 기본은 흑백(모노톤)입니다.
    """
    img = None
    if _BG_PROVIDER:
        try:
            b = _BG_PROVIDER(str(seed), theme, size)
            if b:
                cand = Image.open(io.BytesIO(b)).convert("RGB").resize(size, Image.LANCZOS)
                # 실제 사진인지(= 윤곽이 있는지) 살펴본다.
                # 밋밋한 그라데이션이 돌아오면 직접 그린 장면을 쓴다.
                try:
                    e = ImageOps.grayscale(cand.resize((240, 160))).filter(
                        ImageFilter.FIND_EDGES)
                    energy = sum(e.getdata()) / (240 * 160)
                except Exception:
                    energy = 99
                if energy >= 3.2:
                    img = cand
        except Exception:
            img = None
    if img is None:
        img = _procedural_bg(seed, size)

    if mono:
        g = ImageOps.grayscale(img)
        g = ImageOps.autocontrast(g, cutoff=2)
        img = Image.merge("RGB", (g, g, g))
        # 아주 옅은 잉크색(따뜻한 회색)
        img = Image.blend(img, Image.new("RGB", size, (26, 24, 22)), 0.06)

    # 비네팅 — 인쇄 사진 느낌
    W, H = size
    vig = Image.new("L", size, 0)
    dv = ImageDraw.Draw(vig)
    dv.ellipse([-W * 0.18, -H * 0.18, W * 1.18, H * 1.18], fill=255)
    vig = vig.filter(ImageFilter.GaussianBlur(W / 14))
    img = Image.composite(img, Image.new("RGB", size, (18, 18, 18)), vig)

    # 인쇄 망점 느낌의 아주 옅은 노이즈
    rnd = random.Random(str(seed) + "n")
    noise = Image.new("L", (W // 3, H // 3))
    noise.putdata([rnd.randint(118, 138) for _ in range((W // 3) * (H // 3))])
    noise = noise.resize(size, Image.BILINEAR)
    img = Image.blend(img, Image.merge("RGB", (noise, noise, noise)), 0.06)
    return img


def make_cartoon(title, left_line, right_line, caption, sign="시사주간뉴스 만평",
                 size=(1400, 1000), seed="cartoon"):
    """
    한 컷 만평. 종이질감 위에 잉크선으로 직접 그립니다.
    """
    W, H = size
    rnd = random.Random(str(seed))
    PAPER = (246, 242, 231)
    INK = (26, 26, 28)
    img = Image.new("RGB", (W, H), PAPER)
    d = ImageDraw.Draw(img)

    # 종이 질감
    for _ in range(2200):
        x, y = rnd.randint(0, W - 1), rnd.randint(0, H - 1)
        v = rnd.randint(214, 236)
        d.point((x, y), fill=(v, v - 3, v - 12))

    # 테두리 (겹선)
    d.rectangle([10, 10, W - 11, H - 11], outline=INK, width=5)
    d.rectangle([22, 22, W - 23, H - 23], outline=INK, width=2)

    pad = 44
    # ── 제목
    ft = _pil_font(54, True)
    tl = wrap_pil(d, title or "만평", ft, W - pad * 2)[:1]
    ty = 44
    for ln in tl:
        tw = d.textlength(ln, font=ft)
        d.text(((W - tw) / 2, ty), ln, font=ft, fill=INK)
        ty += 66
    d.line([(pad + 60, ty + 4), (W - pad - 60, ty + 4)], fill=INK, width=3)

    # ── 무대
    stage_top = ty + 40
    stage_bot = H - 190
    ground = stage_bot - 20
    d.line([(pad + 30, ground), (W - pad - 30, ground)], fill=INK, width=4)
    # 바닥 빗금
    for x in range(pad + 34, W - pad - 30, 26):
        d.line([(x, ground), (x - 16, ground + 20)], fill=(90, 88, 86), width=2)

    def figure(cx, base_y, scale, mood="calm", flip=False):
        """아주 단순한 잉크 인물."""
        s = scale
        head_r = int(46 * s)
        head_c = (cx, base_y - int(250 * s))
        # 몸통
        bw, bh = int(86 * s), int(170 * s)
        body = [cx - bw // 2, head_c[1] + head_r - int(6 * s),
                cx + bw // 2, base_y]
        d.polygon([(body[0] + int(14 * s), body[1]), (body[2] - int(14 * s), body[1]),
                   (body[2], body[3]), (body[0], body[3])], fill=(252, 250, 244), outline=INK)
        d.line([(body[0] + int(14 * s), body[1]), (body[0], body[3])], fill=INK, width=4)
        d.line([(body[2] - int(14 * s), body[1]), (body[2], body[3])], fill=INK, width=4)
        d.line([(body[0], body[3]), (body[2], body[3])], fill=INK, width=4)
        # 팔
        ax = -1 if not flip else 1
        d.line([(body[0] + int(8 * s), body[1] + int(26 * s)),
                (cx + ax * int(96 * s), body[1] + int(96 * s))], fill=INK, width=5)
        d.line([(body[2] - int(8 * s), body[1] + int(26 * s)),
                (cx - ax * int(30 * s), body[1] + int(120 * s))], fill=INK, width=5)
        # 머리
        d.ellipse([head_c[0] - head_r, head_c[1] - head_r,
                   head_c[0] + head_r, head_c[1] + head_r],
                  fill=(252, 250, 244), outline=INK, width=5)
        # 눈
        ey = head_c[1] - int(8 * s)
        er = max(3, int(5 * s))
        for ox in (-int(17 * s), int(17 * s)):
            d.ellipse([head_c[0] + ox - er, ey - er, head_c[0] + ox + er, ey + er], fill=INK)
        # 눈썹 · 입
        if mood == "worry":
            d.line([(head_c[0] - int(28 * s), ey - int(20 * s)),
                    (head_c[0] - int(8 * s), ey - int(12 * s))], fill=INK, width=4)
            d.line([(head_c[0] + int(28 * s), ey - int(20 * s)),
                    (head_c[0] + int(8 * s), ey - int(12 * s))], fill=INK, width=4)
            d.arc([head_c[0] - int(22 * s), head_c[1] + int(14 * s),
                   head_c[0] + int(22 * s), head_c[1] + int(44 * s)], 200, 340, fill=INK, width=4)
        else:
            d.line([(head_c[0] - int(28 * s), ey - int(16 * s)),
                    (head_c[0] - int(8 * s), ey - int(20 * s))], fill=INK, width=4)
            d.line([(head_c[0] + int(28 * s), ey - int(16 * s)),
                    (head_c[0] + int(8 * s), ey - int(20 * s))], fill=INK, width=4)
            d.arc([head_c[0] - int(22 * s), head_c[1] + int(4 * s),
                   head_c[0] + int(22 * s), head_c[1] + int(36 * s)], 20, 160, fill=INK, width=4)
        # 머리카락 빗금
        for k in range(6):
            a = math.radians(200 + k * 14)
            d.line([(head_c[0] + math.cos(a) * head_r * 0.95,
                     head_c[1] + math.sin(a) * head_r * 0.95),
                    (head_c[0] + math.cos(a) * head_r * 1.28,
                     head_c[1] + math.sin(a) * head_r * 1.28)], fill=INK, width=3)
        return head_c, head_r

    lh, lr = figure(int(W * 0.29), ground, 1.05, "calm", flip=False)
    rh, rr = figure(int(W * 0.71), ground, 0.92, "worry", flip=True)

    def balloon(anchor, text, box, tail_to):
        if not text:
            return
        x1, y1, x2, y2 = box
        fb = _pil_font(30, False)
        lines = wrap_pil(d, text, fb, (x2 - x1) - 44)[:4]
        need = 26 + len(lines) * 40
        y1 = y2 - max(need, 80)
        d.rounded_rectangle([x1, y1, x2, y2], radius=28,
                            fill=(255, 255, 252), outline=INK, width=4)
        tx = (x1 + x2) / 2
        d.polygon([(tx - 18, y2 - 4), (tx + 18, y2 - 4), tail_to], fill=(255, 255, 252),
                  outline=INK)
        d.line([(tx - 18, y2 - 2), tail_to], fill=INK, width=4)
        d.line([(tx + 18, y2 - 2), tail_to], fill=INK, width=4)
        d.line([(tx - 17, y2 - 3), (tx + 17, y2 - 3)], fill=(255, 255, 252), width=6)
        yy = y1 + 16
        for ln in lines:
            d.text((x1 + 22, yy), ln, font=fb, fill=INK)
            yy += 40

    balloon(lh, left_line, (int(W * 0.06), stage_top, int(W * 0.46), lh[1] - lr - 34),
            (lh[0] + 6, lh[1] - lr - 6))
    balloon(rh, right_line, (int(W * 0.54), stage_top, int(W * 0.94), rh[1] - rr - 34),
            (rh[0] - 6, rh[1] - rr - 6))

    # ── 설명(캡션) 띠
    cy = H - 172
    d.line([(pad, cy), (W - pad, cy)], fill=INK, width=3)
    fc = _pil_font(29, False)
    cl = wrap_pil(d, caption or "", fc, W - pad * 2)[:2]
    yy = cy + 16
    for ln in cl:
        d.text((pad + 2, yy), ln, font=fc, fill=(40, 40, 42))
        yy += 38

    fs = _pil_font(23, True)
    d.text((W - pad - d.textlength(sign, font=fs), H - 62), sign,
           font=fs, fill=(108, 106, 104))
    return img


def make_diagram(title, center, nodes, size=(1200, 900), seed="dg"):
    """가운데 개념 + 둘레 항목의 도해."""
    W, H = size
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    INK = (22, 22, 26)
    d.rectangle([0, 0, W - 1, H - 1], outline=(30, 30, 34), width=3)

    ft = _pil_font(40, True)
    tl = wrap_pil(d, title or "도해", ft, W - 80)[:1]
    yy = 26
    for ln in tl:
        d.text((40, yy), ln, font=ft, fill=INK)
        yy += 52
    d.line([(40, yy + 4), (W - 40, yy + 4)], fill=INK, width=3)

    cx, cy = W // 2, int(H * 0.56)
    R = int(min(W, H) * 0.155)
    palette = [(124, 58, 237), (14, 165, 233), (22, 163, 74), (217, 119, 6),
               (220, 38, 38), (13, 148, 136), (30, 58, 138), (219, 39, 119)]

    nodes = [n for n in (nodes or []) if (n.get("label") or "").strip()][:6]
    n = len(nodes) or 1
    ring = int(min(W, H) * 0.36)
    fl = _pil_font(27, True)
    fd = _pil_font(22, False)

    for i, nd in enumerate(nodes):
        ang = -math.pi / 2 + i * (2 * math.pi / n)
        nx, ny = cx + math.cos(ang) * ring, cy + math.sin(ang) * ring * 0.86
        col = palette[i % len(palette)]
        # 화살표
        sx, sy = cx + math.cos(ang) * (R + 8), cy + math.sin(ang) * (R + 8) * 0.9
        ex, ey = nx - math.cos(ang) * 74, ny - math.sin(ang) * 46
        d.line([(sx, sy), (ex, ey)], fill=col, width=5)
        ah = 16
        d.polygon([(ex, ey),
                   (ex - math.cos(ang - 0.4) * ah, ey - math.sin(ang - 0.4) * ah),
                   (ex - math.cos(ang + 0.4) * ah, ey - math.sin(ang + 0.4) * ah)], fill=col)
        # 상자
        bw, bh = 250, 108
        bx1, by1 = nx - bw / 2, ny - bh / 2
        d.rounded_rectangle([bx1, by1, bx1 + bw, by1 + bh], radius=16,
                            fill=(248, 250, 255), outline=col, width=4)
        lab = wrap_pil(d, nd.get("label", ""), fl, bw - 24)[:1]
        dsc = wrap_pil(d, nd.get("desc", ""), fd, bw - 24)[:2]
        ty = by1 + 12
        for ln in lab:
            d.text((bx1 + 12, ty), ln, font=fl, fill=col)
            ty += 34
        for ln in dsc:
            d.text((bx1 + 12, ty), ln, font=fd, fill=(70, 74, 88))
            ty += 28

    # 가운데
    d.ellipse([cx - R, cy - R * 0.9, cx + R, cy + R * 0.9],
              fill=(17, 24, 39), outline=(17, 24, 39))
    fcn = _pil_font(30, True)
    cl = wrap_pil(d, center or "이번 주", fcn, R * 1.7)[:3]
    ty = cy - len(cl) * 19
    for ln in cl:
        d.text((cx - d.textlength(ln, font=fcn) / 2, ty), ln, font=fcn, fill=(255, 255, 255))
        ty += 38
    return img


def _pil_to_reader(img):
    b = io.BytesIO()
    img.save(b, format="PNG")
    b.seek(0)
    return ImageReader(b)


def pil_to_bytes(img, fmt="PNG", quality=88):
    b = io.BytesIO()
    if fmt.upper() in ("JPG", "JPEG"):
        img.convert("RGB").save(b, format="JPEG", quality=quality, optimize=True)
    else:
        img.save(b, format="PNG", optimize=True)
    return b.getvalue()


# ==============================================================================
# 3. 지면 설계 (내용 → 8면 구성)
# ==============================================================================
PAGE_NAMES = ["종합", "이슈 해설", "사회·시사", "국제·경제",
              "만평·사설", "도해·통계", "문화·현장", "신앙·목회"]


def _fallback_article(name, items):
    """AI 가 없어도 기사 모양이 나오도록 헤드라인만으로 만든 기사."""
    if not items:
        return None
    top = items[0]
    body = []
    if len(items) > 1:
        body.append("같은 기간 " + name + " 분야에서는 "
                    + ", ".join(_clean(it["title"])[:36] for it in items[1:4])
                    + " 등의 보도가 이어졌다.")
    body.append("이상은 지난 한 주간 실제 보도된 기사 제목을 모은 것으로, "
                "자세한 내용은 각 매체의 원문을 확인해야 한다.")
    return {
        "section": name,
        "headline": _clean(top["title"]),
        "lead": f"{_clean(top.get('source','')) or '주요 매체'} 등이 지난 한 주간 보도한 내용이다.",
        "body": body,
        "source": _clean(top.get("source", "")),
    }


def build_spec(paper_title, publisher, win_s, win_e, results, ai_json=None,
               ai_text="", keywords=None, chart_png=None, issue_no="",
               page_sections=None, theme="도시 · 일상", seed="np",
               editor="", church=""):
    """
    results      : {섹션이름: {"group","color","query","items":[{title,link,source,when}]}}
    ai_json      : AI 가 만든 기사화 결과(dict) — 없으면 헤드라인만으로 만듭니다
    page_sections: {3:[섹션이름...], 4:[...], 7:[...]}  면별 배치
    """
    aj = ai_json or {}
    kws = keywords or []
    names = [n for n in results.keys() if results[n]["items"]]

    if not page_sections:
        page_sections = {3: [], 4: [], 7: []}
        for i, n in enumerate(names):
            page_sections[[3, 4, 7][i % 3]].append(n)

    # ── 1면 톱
    top = aj.get("top") or {}
    first_items = []
    for n in names:
        first_items.extend(results[n]["items"][:2])
    if not top.get("headline") and first_items:
        fb = _fallback_article(names[0] if names else "종합", first_items)
        top = {"kicker": "이번 주 톱", "headline": fb["headline"],
               "sub": "", "lead": fb["lead"], "body": fb["body"],
               "photo_caption": "지난 한 주간의 현장. (자료 이미지)"}
    top.setdefault("kicker", "이번 주 톱")
    top.setdefault("headline", "이번 주 시사주간뉴스")
    top.setdefault("body", [])
    top.setdefault("lead", "")
    top.setdefault("photo_caption", "이번 주 현장. (자료 이미지)")

    second = aj.get("second") or []
    if not second:
        second = []
        for n in names[1:3]:
            a = _fallback_article(n, results[n]["items"])
            if a:
                second.append({"headline": a["headline"], "body": a["body"][:1],
                               "source": a["source"]})

    flow3 = aj.get("flow3") or []
    if not flow3 and ai_text:
        # AI 해설 본문에서 '큰 흐름' 부분을 긁어 온다
        blk = re.split(r"\n(?=[🔎📋🙏⚠️🗞️])", ai_text)
        for b in blk:
            if "큰 흐름" in b.splitlines()[0] if b.splitlines() else False:
                for m in re.findall(r"-\s*\d+\.\s*(.+)", b):
                    t = _clean(m)
                    head, _, rest = t.partition("—")
                    flow3.append({"title": _clean(head) or t[:24],
                                  "body": [_clean(rest) or t]})
    if not flow3:
        for w, c in kws[:3]:
            rel = []
            for n in names:
                for it in results[n]["items"]:
                    if w in it["title"]:
                        rel.append(_clean(it["title"]))
            flow3.append({"title": f"‘{w}’ — 이번 주 {c}건",
                          "body": [("관련 보도: " + " / ".join(rel[:4])) if rel else
                                   "이번 주 자주 등장한 낱말이다."]})

    cart = aj.get("cartoon") or {}
    cart.setdefault("title", "이번 주 만평")
    cart.setdefault("left_line", "세상은 늘 바쁘게 돌아갑니다.")
    cart.setdefault("right_line", "그 속에서 우리는 어디를 보고 있습니까?")
    cart.setdefault("caption", "한 주간의 뉴스를 지나며, 교회는 무엇을 붙들어야 하는가.")

    edit = aj.get("editorial") or {}
    edit.setdefault("title", "사설 — 소란한 한 주 끝에서")
    edit.setdefault("body", [
        "한 주간의 기사를 모아 놓고 보면 세상의 소리가 얼마나 큰지 알게 된다.",
        "그러나 교회의 자리는 그 소리를 따라가는 자리가 아니라, "
        "그 소리 속에서 하나님의 뜻을 분별하는 자리다.",
        "이번 주에도 우리는 말씀 앞에 다시 선다."])

    tbl = aj.get("table") or {}
    if not tbl.get("rows"):
        rows = [[n, str(len(results[n]["items"])), results[n]["group"]]
                for n in names]
        rows.sort(key=lambda r: -int(r[1]))
        tbl = {"title": "섹션별 수집 기사 수",
               "headers": ["섹션", "기사 수", "분류"], "rows": rows[:14]}

    dgm = aj.get("diagram") or {}
    if not dgm.get("nodes"):
        dgm = {"title": "이번 주 흐름 한눈에",
               "center": "이번 주\n한국·세계",
               "nodes": [{"label": w, "desc": f"관련 보도 {c}건"} for w, c in kws[:5]]}

    links = aj.get("sermon_links") or []
    if not links:
        links = [{"event": _clean(it["title"])[:40], "text": "—",
                  "use": "강단에서 다룰 때는 사실관계를 먼저 확인하십시오."}
                 for n in names[:5] for it in results[n]["items"][:1]][:5]

    prayer = aj.get("prayer") or [
        "혼란한 세상 가운데 진리를 분별하게 하옵소서.",
        "고통 받는 이웃과 재난의 자리를 기억하게 하옵소서.",
        "이 나라의 위정자들에게 지혜를 주옵소서.",
        "교회가 세상의 소리보다 말씀에 먼저 귀 기울이게 하옵소서."]

    total = sum(len(v["items"]) for v in results.values())
    spec = {
        "title": paper_title or "시사주간뉴스",
        "sub": aj.get("masthead_sub") or "한 주간의 세상을 말씀의 눈으로 읽습니다",
        "publisher": publisher or "화광교회",
        "editor": editor,
        "church": church or publisher or "화광교회",
        "issue": issue_no or f"제 {win_e.isocalendar()[1]} 호",
        "date_line": f"{win_e.year}년 {win_e.month}월 {win_e.day}일 "
                     f"{'월화수목금토일'[win_e.weekday()]}요일",
        "period": f"{win_s.strftime('%Y.%m.%d')} ~ {win_e.strftime('%Y.%m.%d')}",
        "total": total,
        "nsec": len(names),
        "keywords": kws,
        "results": results,
        "names": names,
        "page_sections": page_sections,
        "top": top, "second": second, "flow3": flow3,
        "cartoon": cart, "editorial": edit, "table": tbl, "diagram": dgm,
        "sermon_links": links, "prayer": prayer,
        "ai_text": ai_text,
        "chart_png": chart_png,
        "theme": theme,
        "seed": seed,
    }
    return spec


# ==============================================================================
# 2-1. 구글 AI(Gemini) 사진 만들기
#   · 추가 설치 없이 표준 라이브러리(urllib)로 Gemini REST API 를 부릅니다.
#   · 목사님 API 키로 쓸 수 있는 '이미지 모델'을 스스로 찾아 차례로 시도합니다.
#   · 한 장이라도 실패하면 그 자리는 기본 그림으로 채웁니다(신문은 항상 완성).
# ==============================================================================
import json as _json
import base64 as _b64
import time as _time
import urllib.request as _ureq
import urllib.error as _uerr

GEMINI_BASE = "https://generativelanguage.googleapis.com/v1beta"
# 먼저 시도할 순서 (목록에 없는 이름이면 자동으로 건너뜁니다)
IMAGE_MODEL_PREF = [
    "gemini-3.1-flash-image", "gemini-3.1-flash-image-preview",
    "gemini-2.5-flash-image", "gemini-2.5-flash-image-preview",
    "gemini-3-pro-image", "gemini-3-pro-image-preview",
    "gemini-2.0-flash-preview-image-generation",
    "gemini-2.0-flash-exp-image-generation",
]
_IMG_STATE = {"ok_model": None, "bad": set(), "models": None, "models_at": 0}

PHOTO_LABEL_AI = "(구글 AI 생성 이미지 · 실제 현장 사진 아님)"
PHOTO_LABEL_BASIC = "(자료 이미지)"


def _http_json(url, api_key, body=None, timeout=60):
    data = None if body is None else _json.dumps(body).encode("utf-8")
    req = _ureq.Request(url, data=data, method="POST" if body is not None else "GET",
                        headers={"x-goog-api-key": api_key,
                                 "Content-Type": "application/json"})
    with _ureq.urlopen(req, timeout=timeout) as r:
        return _json.loads(r.read().decode("utf-8"))


def list_image_models(api_key, timeout=15):
    """이 API 키로 쓸 수 있는 이미지 생성 모델 이름들."""
    if _IMG_STATE["models"] is not None and _time.time() - _IMG_STATE["models_at"] < 1800:
        return _IMG_STATE["models"]
    found = []
    try:
        url = f"{GEMINI_BASE}/models?pageSize=200"
        for _ in range(5):
            js = _http_json(url, api_key, timeout=timeout)
            for m in js.get("models", []):
                name = m.get("name", "").split("/")[-1]
                meths = m.get("supportedGenerationMethods", [])
                if "image" in name and "generateContent" in meths and not name.startswith("imagen"):
                    found.append(name)
            tok = js.get("nextPageToken")
            if not tok:
                break
            url = f"{GEMINI_BASE}/models?pageSize=200&pageToken={tok}"
    except Exception:
        found = []
    pref = [m for m in IMAGE_MODEL_PREF if m in found]
    rest = sorted([m for m in found if m not in pref], reverse=True)
    models = pref + rest
    if not models:                       # 목록을 못 받으면 알려진 이름으로 직접 시도
        models = list(IMAGE_MODEL_PREF)
    _IMG_STATE["models"] = models
    _IMG_STATE["models_at"] = _time.time()
    return models


def _extract_image(js):
    for cand in js.get("candidates", []) or []:
        for part in (cand.get("content") or {}).get("parts", []) or []:
            inl = part.get("inlineData") or part.get("inline_data")
            if inl and inl.get("data"):
                return _b64.b64decode(inl["data"])
    return None


def google_image(api_key, prompt, aspect="16:9", timeout=90):
    """
    Gemini 로 그림 한 장. 성공하면 (bytes, 모델이름), 실패하면 (None, 사유).
    """
    if not api_key:
        return None, "API 키가 없습니다"
    deadline = _time.time() + timeout
    models = list_image_models(api_key)
    if _IMG_STATE["ok_model"] in models:
        models = [_IMG_STATE["ok_model"]] + [m for m in models if m != _IMG_STATE["ok_model"]]
    last = "쓸 수 있는 이미지 모델이 없습니다"
    for m in models:
        if m in _IMG_STATE["bad"]:
            continue
        left = deadline - _time.time()
        if left < 8:
            return None, last + " · 제한 시간 초과"
        url = f"{GEMINI_BASE}/models/{m}:generateContent"
        bodies = [
            {"contents": [{"parts": [{"text": prompt}]}],
             "generationConfig": {"responseModalities": ["TEXT", "IMAGE"],
                                  "imageConfig": {"aspectRatio": aspect}}},
            {"contents": [{"parts": [{"text": prompt}]}],
             "generationConfig": {"responseModalities": ["TEXT", "IMAGE"]}},
        ]
        for body in bodies:
            try:
                js = _http_json(url, api_key, body,
                                timeout=max(8, min(timeout, deadline - _time.time())))
                b = _extract_image(js)
                if b:
                    _IMG_STATE["ok_model"] = m
                    return b, m
                fb = (js.get("promptFeedback") or {}).get("blockReason")
                last = f"{m}: 그림이 오지 않았습니다" + (f" (차단: {fb})" if fb else "")
                break                                   # 같은 모델 재시도 불필요
            except _uerr.HTTPError as e:
                try:
                    msg = e.read().decode("utf-8", "ignore")[:300]
                except Exception:
                    msg = ""
                if e.code == 429:
                    return None, "구글 이미지 무료 한도(분당/일일)를 넘었습니다 (429)"
                if "API_KEY_INVALID" in msg or "API key not valid" in msg:
                    return None, "API 키가 올바르지 않습니다 (사이드바 AI 연결 설정 확인)"
                if e.code == 400 and body is bodies[0]:
                    last = f"{m}: 400 {msg[:120]}"
                    continue                            # imageConfig 없이 다시
                if e.code in (400, 403, 404):
                    _IMG_STATE["bad"].add(m)
                last = f"{m}: HTTP {e.code} {msg[:120]}"
                break
            except (_uerr.URLError, TimeoutError, OSError) as e:
                # 인터넷 연결 자체가 안 되면 다른 모델도 똑같다 — 바로 멈춘다
                return None, f"구글 서버에 연결하지 못했습니다 ({type(e).__name__}: {str(e)[:80]})"
            except Exception as e:
                last = f"{m}: {type(e).__name__} {str(e)[:100]}"
                break
    return None, last


def photo_prompt(subject, kind="news"):
    """신문 사진용 지시문. 실존 인물·글자·선정적 장면을 피하게 합니다."""
    subject = _clean(subject)[:160]
    if kind == "church":
        scene = ("따뜻한 빛이 드는 한국 교회 예배당에서 조용히 기도하는 성도들의 "
                 "뒷모습, 성경책과 창문 빛")
        return (f"{scene}. 사실적인 다큐멘터리 보도사진 스타일, 35mm 렌즈, 자연광, "
                "가로로 긴 구도. 얼굴이 식별되지 않게 뒷모습·원경으로. "
                "글자·로고·워터마크·자막은 절대 넣지 마세요.")
    return (f"다음 한국 신문 기사에 어울리는 보도사진 한 장을 만들어 주세요. 기사 제목: 「{subject}」. "
            "조건: 사실적인 다큐멘터리 보도사진 스타일, 자연광, 35mm 렌즈 느낌, "
            "한국의 실제 거리·건물·풍경처럼 보이게, 가로로 긴 구도. "
            "정치인·연예인 등 실존 인물을 그리지 말고, 사람은 얼굴이 식별되지 않게 "
            "뒷모습이나 원경으로만. 사건·사고·재난 기사라면 참혹한 장면 대신 "
            "상징적인 장면(빈 거리, 구호 물품, 흐린 하늘 등)으로. "
            "글자·간판 문구·로고·워터마크·자막은 절대 넣지 마세요.")


def photo_subjects(spec):
    """각 사진 자리에 들어갈 기사 제목."""
    subj = {"top": (spec["top"].get("headline", "") + " — " + spec["top"].get("sub", "")).strip(" —")}
    for no, key in ((3, "p3"), (4, "p4"), (7, "p7")):
        names = [n for n in spec["page_sections"].get(no, [])
                 if spec["results"].get(n, {}).get("items")]
        if names:
            subj[key] = _clean(spec["results"][names[0]]["items"][0]["title"])
        else:
            subj[key] = spec["top"].get("headline", "이번 주 한국 사회")
    subj["p8"] = ""
    return subj


def generate_google_photos(spec, api_key, progress=None, budget=150, workers=3):
    """
    사진 5장(1면·3면·4면·7면·8면)을 구글 AI로 만들어 spec["photo_bytes"] 에 담습니다.
    반환: (성공 장수, 오류 목록, 사용한 모델)
    """
    from concurrent.futures import ThreadPoolExecutor, as_completed
    subj = photo_subjects(spec)
    jobs = {k: photo_prompt(v, "church" if k == "p8" else "news") for k, v in subj.items()}
    out, errs, used = {}, [], ""
    t0 = _time.time()

    # 첫 장은 혼자 먼저 — 모델 찾기와 한도 확인을 한 번에 끝낸다
    first = "top"
    b, info = google_image(api_key, jobs[first], "16:9",
                           timeout=min(90, max(20, budget - 5)))
    if progress:
        try:
            progress(1 / len(jobs), first)
        except Exception:
            pass
    if b:
        out[first] = b
        used = info
    else:
        errs.append(f"1면 사진: {info}")
        spec["photo_bytes"] = out
        return 0, errs, used                    # 첫 장부터 안 되면 나머지도 안 된다

    rest = [k for k in jobs if k != first]
    ex = ThreadPoolExecutor(max_workers=workers)
    futs = {ex.submit(google_image, api_key, jobs[k], "16:9",
                      max(20, min(90, budget - (_time.time() - t0)))): k for k in rest}
    done = 1
    try:
        for f in as_completed(futs, timeout=max(10, budget - (_time.time() - t0))):
            k = futs[f]
            try:
                b, info = f.result(timeout=1)
            except Exception as e:
                b, info = None, f"{type(e).__name__}"
            if b:
                out[k] = b
                used = used or info
            else:
                errs.append(f"{k} 사진: {info}")
            done += 1
            if progress:
                try:
                    progress(done / len(jobs), k)
                except Exception:
                    pass
    except Exception:
        errs.append(f"제한 시간({budget}초) 안에 끝나지 않은 사진은 기본 그림으로 채웠습니다")
    try:
        ex.shutdown(wait=False, cancel_futures=True)
    except TypeError:
        ex.shutdown(wait=False)
    spec["photo_bytes"] = out
    return len(out), errs, used


def photo_caption(spec, key, text):
    """사진 설명 끝에 'AI 생성' 또는 '자료 이미지' 표시를 정확히 붙입니다."""
    is_ai = key in (spec.get("photo_bytes") or {})
    label = PHOTO_LABEL_AI if is_ai else PHOTO_LABEL_BASIC
    t = _clean(text).replace("(자료 이미지)", "").replace(PHOTO_LABEL_AI, "").strip()
    return f"{t} {label}".strip()


def _photo_from_bytes(b, size, mono):
    img = Image.open(io.BytesIO(b)).convert("RGB")
    if mono:
        g = ImageOps.autocontrast(ImageOps.grayscale(img), cutoff=1)
        img = Image.merge("RGB", (g, g, g))
    # 너무 크면 줄여서 PDF 용량을 아낀다
    if img.size[0] > 1600:
        img = img.resize((1600, int(img.size[1] * 1600 / img.size[0])), Image.LANCZOS)
    return img


def render_images(spec, progress=None):
    """지면에 들어갈 그림을 미리 만들어 spec 에 담습니다."""
    seed = spec["seed"]
    th = spec["theme"]
    pb = spec.get("photo_bytes") or {}
    mono = bool(spec.get("photo_mono", True))
    imgs = {}

    def photo(key, size):
        if key in pb:
            try:
                return _photo_from_bytes(pb[key], size, mono)
            except Exception:
                pb.pop(key, None)
        return make_photo(f"{seed}|{key}", size, th, mono=mono)

    steps = [
        ("top", lambda: photo("top", (1400, 900))),
        ("p3", lambda: photo("p3", (1200, 780))),
        ("p4", lambda: photo("p4", (1200, 780))),
        ("p7", lambda: photo("p7", (1200, 780))),
        ("cartoon", lambda: make_cartoon(spec["cartoon"]["title"],
                                         spec["cartoon"]["left_line"],
                                         spec["cartoon"]["right_line"],
                                         spec["cartoon"]["caption"],
                                         sign=f"만평 · {spec['title']}",
                                         seed=f"{seed}|c")),
        ("diagram", lambda: make_diagram(spec["diagram"]["title"],
                                         spec["diagram"].get("center", "이번 주"),
                                         spec["diagram"].get("nodes", []),
                                         seed=f"{seed}|d")),
        ("p8", lambda: photo("p8", (1200, 700))),
    ]
    for i, (k, fn) in enumerate(steps):
        if progress:
            try:
                progress((i + 1) / len(steps), k)
            except Exception:
                pass
        try:
            imgs[k] = fn()
        except Exception:
            imgs[k] = Image.new("RGB", (1200, 780), (230, 230, 232))
    if spec.get("chart_png"):
        try:
            imgs["chart"] = Image.open(io.BytesIO(spec["chart_png"])).convert("RGB")
        except Exception:
            pass
    spec["images"] = imgs
    return spec


# ==============================================================================
# 4. PDF 조판 (A3 세로 · 8면)
# ==============================================================================
PW, PH = A3                       # 841.89 x 1190.55 pt  (297 x 420 mm)
ML = MR = 13 * mm
MT = 11 * mm
MB = 12 * mm
CW = PW - ML - MR                 # 본문 폭
NCOL = 4
GUT = 4.6 * mm
COLW = (CW - GUT * (NCOL - 1)) / NCOL

INK = rlcolors.HexColor("#141414")
GREY = rlcolors.HexColor("#5a5a5e")
LINE = rlcolors.HexColor("#1d1d1f")
SOFT = rlcolors.HexColor("#9a9a9e")


def colx(i):
    return ML + i * (COLW + GUT)


def spanw(n):
    return COLW * n + GUT * (n - 1)


def _style(name, size, leading=None, font=None, align=TA_JUSTIFY, color=INK,
           space_after=3.2, first_indent=0):
    return ParagraphStyle(
        name, fontName=font or F_REG, fontSize=size,
        leading=leading or size * 1.52, alignment=align,
        textColor=color, wordWrap="CJK", spaceAfter=space_after,
        firstLineIndent=first_indent)


def draw_rule(c, x1, y, x2, w=0.7, color=LINE):
    c.setStrokeColor(color)
    c.setLineWidth(w)
    c.line(x1, y, x2, y)


def draw_vrule(c, x, y1, y2, w=0.4, color=SOFT):
    c.setStrokeColor(color)
    c.setLineWidth(w)
    c.line(x, y1, x, y2)


def draw_headline(c, text, x, ytop, w, max_size=40, min_size=15, max_lines=3,
                  font=None, color=INK, lead=1.16, align="left"):
    """자동으로 크기를 줄여 가며 제목을 앉힌다. 마지막 y(아래쪽)를 돌려준다."""
    font = font or F_BOLD
    text = _clean(text)
    size, lines = None, []
    floor = max(min_size, max_size * 0.74)      # 여기까지는 줄 수를 줄이는 쪽이 낫다
    for s in range(int(max_size), int(min_size) - 1, -1):
        ls = wrap_pdf(text, font, s, w)
        if len(ls) > max_lines:
            continue
        if size is None:
            size, lines = s, ls
        elif len(ls) < len(lines) and s >= floor:
            size, lines = s, ls
        elif s < floor:
            break
    if size is None:
        size = min_size
        lines = wrap_pdf(text, font, size, w)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1][:-1] + "…"
    c.setFont(font, size)
    c.setFillColor(color)
    y = ytop - size
    for ln in lines:
        if align == "center":
            c.drawCentredString(x + w / 2, y, ln)
        else:
            c.drawString(x, y, ln)
        y -= size * lead
    return y + size * lead - size * 0.28


def draw_label(c, text, x, y, color=rlcolors.HexColor("#111827"),
               fg=rlcolors.white, size=8.4, padx=5, pady=3):
    """섹션 이름표(까만 바탕 흰 글씨)"""
    tw = sw(text, F_BOLD, size)
    c.setFillColor(color)
    c.rect(x, y - pady, tw + padx * 2, size + pady * 1.9, stroke=0, fill=1)
    c.setFillColor(fg)
    c.setFont(F_BOLD, size)
    c.drawString(x + padx, y + 0.6, text)
    return tw + padx * 2


def draw_kicker(c, text, x, y, color=rlcolors.HexColor("#b91c1c"), size=9.6):
    c.setFillColor(color)
    c.setFont(F_BOLD, size)
    c.drawString(x, y, text)
    return y - size * 1.3


def place_image(c, img, x, y, w, h, caption=None, cap_size=7.8):
    """가운데를 잘라 채워 넣기(커버) + 캡션"""
    if img is None:
        return y
    iw, ih = img.size
    tr, br = w / h, iw / ih
    if br > tr:
        nw = int(ih * tr)
        box = ((iw - nw) // 2, 0, (iw + nw) // 2, ih)
    else:
        nh = int(iw / tr)
        box = (0, (ih - nh) // 2, iw, (ih + nh) // 2)
    im = img.crop(box)
    c.drawImage(_pil_to_reader(im), x, y, w, h, mask=None)
    c.setStrokeColor(rlcolors.HexColor("#2a2a2c"))
    c.setLineWidth(0.6)
    c.rect(x, y, w, h, stroke=1, fill=0)
    if caption:
        lines = wrap_pdf(caption, F_REG, cap_size, w)[:2]
        yy = y - cap_size - 2.4
        c.setFillColor(GREY)
        c.setFont(F_REG, cap_size)
        for ln in lines:
            c.drawString(x, yy, ln)
            yy -= cap_size * 1.28
        return yy
    return y


def flow(c, flowables, rects):
    """여러 칸(rect)에 글을 흘려 넣는다. 남은 것을 돌려준다."""
    left = list(flowables)
    for (x, y, w, h) in rects:
        if not left or h <= 6:
            break
        f = Frame(x, y, w, h, leftPadding=0, rightPadding=0,
                  topPadding=0, bottomPadding=0, showBoundary=0)
        try:
            f.addFromList(left, c)
        except Exception:
            break
    return left


def measure(c, flowables, w):
    """흘림 요소들의 전체 높이를 재 본다."""
    tot = 0.0
    for f in flowables:
        try:
            _, h = f.wrapOn(c, w, 100000)
            tot += h
            try:
                tot += float(f.getSpaceAfter() or 0)
            except Exception:
                pass
        except Exception:
            tot += 12
    return tot


def col_rules(c, y_top, y_bot, ncols=NCOL, first=0):
    """단과 단 사이 세로 괘선 — 진짜 신문처럼 보이게 한다."""
    if y_top - y_bot < 24:
        return
    for i in range(first, first + ncols - 1):
        draw_vrule(c, colx(i) + COLW + GUT / 2, y_bot + 2, y_top - 2, 0.35,
                   rlcolors.HexColor("#b9b9bd"))


def flow_balanced(c, flowables, y_top, y_bot, ncols=NCOL, first=0, rules=True):
    """
    글을 ncols 개의 단에 '고르게' 나누어 흘린다.
    (한 단만 꽉 차고 나머지가 비는 현상을 막는다)
    반환: (남은 요소, 실제로 쓴 아래쪽 y)
    """
    flowables = [f for f in flowables if f is not None]
    if not flowables:
        return [], y_top
    avail = max(10, y_top - y_bot)
    total = measure(c, flowables, COLW)
    target = min(avail, max(48, total / max(1, ncols) * 1.05 + 14))
    rects = [(colx(first + i), y_top - target, COLW, target) for i in range(ncols)]
    left = flow(c, flowables, rects)
    used_bot = y_top - target
    if rules:
        col_rules(c, y_top, used_bot, ncols, first)
    return left, used_bot


def art_flowables(headline=None, lead=None, body=None, source=None,
                  hsize=15, bsize=8.9, label=None, label_color=None):
    """기사 한 꼭지를 흘림 요소로."""
    fl = []
    if label:
        fl.append(Paragraph(f'<font color="#b91c1c"><b>{esc(label)}</b></font>',
                            _style("lb", 8.4, 11, F_BOLD, TA_LEFT)))
    if headline:
        fl.append(Paragraph(esc(headline),
                            _style("h", hsize, hsize * 1.26, F_BOLD, TA_LEFT,
                                   space_after=4.4)))
    if lead:
        fl.append(Paragraph(f"<b>{esc(lead)}</b>",
                            _style("ld", bsize + 0.4, (bsize + 0.4) * 1.55, F_BOLD,
                                   space_after=3.6)))
    for p in (body or []):
        p = _clean(p)
        if p:
            fl.append(Paragraph(esc(p), _style("b", bsize, bsize * 1.62, F_REG,
                                               first_indent=bsize)))
    if source:
        fl.append(Paragraph(f'<font color="#6b7280">{esc(source)}</font>',
                            _style("s", 7.8, 10.5, F_REG, TA_LEFT, space_after=7)))
    fl.append(Spacer(1, 5))
    return fl


def brief_flowables(title, items, bsize=8.2):
    fl = [Paragraph(esc(title), _style("bt", 10.6, 14, F_BOLD, TA_LEFT, space_after=4))]
    for i, it in enumerate(items, 1):
        fl.append(Paragraph(f'<b>{i}.</b> {esc(it)}',
                            _style("bi", bsize, bsize * 1.5, F_REG, TA_LEFT,
                                   space_after=2.6)))
    fl.append(Spacer(1, 6))
    return fl


def make_table(headers, rows, w, fsize=8.2, headbg="#111827"):
    data = [[Paragraph(f'<b>{esc(h)}</b>',
                       _style("th", fsize, fsize * 1.4, F_BOLD, TA_CENTER,
                              color=rlcolors.white)) for h in headers]]
    for r in rows:
        data.append([Paragraph(esc(x), _style("td", fsize, fsize * 1.4, F_REG,
                                              TA_CENTER if i else TA_LEFT))
                     for i, x in enumerate(r)])
    n = len(headers)
    if n == 3:
        cw = [w * 0.52, w * 0.2, w * 0.28]
    elif n == 2:
        cw = [w * 0.62, w * 0.38]
    else:
        cw = [w / n] * n
    t = Table(data, colWidths=cw, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rlcolors.HexColor(headbg)),
        ("GRID", (0, 0), (-1, -1), 0.4, rlcolors.HexColor("#9ca3af")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [rlcolors.white, rlcolors.HexColor("#f3f4f6")]),
    ]))
    return t


# ------------------------------------------------------------------ 지면 머리
def masthead(c, spec):
    y = PH - MT
    c.setFillColor(INK)
    c.setFont(F_REG, 8.4)
    c.drawString(ML, y - 9, spec["date_line"])
    c.drawRightString(PW - MR, y - 9, f'{spec["publisher"]} · {spec["issue"]}')
    draw_rule(c, ML, y - 15, PW - MR, 2.0)
    draw_rule(c, ML, y - 18.4, PW - MR, 0.6)

    t = spec["title"]
    size = 76
    while sw(t, F_BOLD, size) > CW * 0.78 and size > 32:
        size -= 1
    c.setFont(F_BOLD, size)
    c.setFillColor(INK)
    ty = y - 30 - size
    c.drawCentredString(PW / 2, ty, t)

    c.setFont(F_REG, 10.6)
    c.setFillColor(GREY)
    c.drawCentredString(PW / 2, ty - size * 0.34 - 13, spec["sub"])

    y2 = ty - size * 0.34 - 24
    draw_rule(c, ML, y2, PW - MR, 0.6)
    draw_rule(c, ML, y2 - 3.4, PW - MR, 2.0)

    c.setFont(F_REG, 8)
    c.setFillColor(GREY)
    c.drawString(ML, y2 - 16, f'수집 기간 {spec["period"]}  ·  기사 {spec["total"]}건  '
                              f'·  섹션 {spec["nsec"]}개  ·  출처 Google 뉴스')
    c.drawRightString(PW - MR, y2 - 16, "A3 세로 · 8면")
    return y2 - 30


def running_head(c, spec, no, name):
    y = PH - MT
    c.setFillColor(INK)
    c.setFont(F_BOLD, 13)
    if no % 2 == 0:
        c.drawString(ML, y - 12, f"{no}")
        c.setFont(F_REG, 9)
        c.drawString(ML + 16, y - 11, f'{name}')
        c.setFont(F_REG, 8.2)
        c.setFillColor(GREY)
        c.drawRightString(PW - MR, y - 11, f'{spec["title"]}  {spec["date_line"]}')
    else:
        c.drawRightString(PW - MR, y - 12, f"{no}")
        c.setFont(F_REG, 9)
        c.drawRightString(PW - MR - 16, y - 11, f'{name}')
        c.setFont(F_REG, 8.2)
        c.setFillColor(GREY)
        c.drawString(ML, y - 11, f'{spec["title"]}  {spec["date_line"]}')
    draw_rule(c, ML, y - 17, PW - MR, 1.4)
    return y - 24


def page_foot(c, spec, no):
    draw_rule(c, ML, MB + 12, PW - MR, 0.5, SOFT)
    c.setFillColor(SOFT)
    c.setFont(F_REG, 7.4)
    c.drawString(ML, MB + 3, f'{spec["title"]} · {spec["period"]}')
    c.drawCentredString(PW / 2, MB + 3, f"- {no} -")
    c.drawRightString(PW - MR, MB + 3, spec["publisher"])


# ------------------------------------------------------------------ 각 면
def _sec_briefs(spec, names, limit=5, bsize=8.0):
    """섹션별 짧은 기사 묶음 — 지면이 비지 않게 채워 준다."""
    fl = []
    for n in names:
        its = spec["results"].get(n, {}).get("items", [])
        if not its:
            continue
        fl += brief_flowables(
            n, [f'{_clean(it["title"])}  ({_clean(it["source"])})' for it in its[:limit]],
            bsize=bsize)
    return fl


FILLERS = [
    ("이 주의 말씀",
     ["“여호와여 주의 말씀이 내 발에 등이요 내 길에 빛이니이다.” (시 119:105)",
      "한 주간의 소식을 다 읽고 나면 마음이 어지럽습니다. "
      "그때 우리가 돌아갈 자리는 언제나 말씀입니다.",
      "이 신문은 세상 소식을 전하기 위해서가 아니라, "
      "그 소식 앞에서 말씀을 다시 펴기 위해 만들어졌습니다."]),
    ("편집실에서",
     ["이 신문의 모든 헤드라인은 지난 한 주간 실제로 보도된 기사에서 가져왔습니다.",
      "기사 원문 링크는 앱의 ‘시사주간뉴스’ 화면에서 그대로 확인하실 수 있습니다.",
      "해설과 사설은 수집된 기사만을 근거로 정리한 것이며, "
      "사실관계가 확정되지 않은 사안은 단정하지 않았습니다."]),
    ("독자와 함께 — 이번 주 나눔 질문",
     ["이번 주 소식 가운데 우리 교회가 함께 기도할 제목은 무엇입니까?",
      "이 소식들 앞에서 성도들이 가장 두려워하는 것은 무엇이겠습니까?",
      "한 사람의 이웃을 위해 이번 주에 할 수 있는 일 한 가지는 무엇입니까?"]),
    ("이 신문을 읽는 법",
     ["1면은 이번 주 가장 큰 흐름, 2면은 해설과 통계입니다.",
      "3·4·7면은 목사님이 직접 고르신 섹션의 기사입니다.",
      "5면은 만평과 사설, 6면은 도해와 표, 8면은 설교·목회 연결입니다."]),
    ("자료 이용 안내",
     ["기사 제목과 매체명은 원문 그대로이며, 요약과 해설은 설교 준비용입니다.",
      "인용하실 때에는 반드시 원문을 확인하신 뒤 매체명을 함께 밝혀 주십시오.",
      "사진은 구글 AI가 기사 내용을 바탕으로 만들었거나 앱이 그린 자료 이미지로, "
      "실제 사건 현장 사진이 아닙니다."]),
]


def filler_box(c, spec, y_top, y_bot, idx=0, ncols=3):
    """
    지면이 남을 때 넣는 고정 꼭지(박스 기사). 내용 높이에 맞춰 상자를 그립니다.
    """
    gap = y_top - y_bot
    if gap < 72:
        return y_bot
    title, paras = FILLERS[idx % len(FILLERS)]
    inner_w = (CW - 28 - GUT * (ncols - 1)) / ncols
    fl = [Paragraph(esc(t), _style("fb", 8.8, 13.6, F_REG, space_after=4))
          for t in paras]
    total = measure(c, fl, inner_w)
    need = min(gap, max(62, total / ncols + 52))
    top, bot = y_top, y_top - need

    c.setStrokeColor(LINE)
    c.setLineWidth(1.1)
    c.rect(ML, bot, CW, need, stroke=1, fill=0)
    c.setLineWidth(0.4)
    c.rect(ML + 3, bot + 3, CW - 6, need - 6, stroke=1, fill=0)

    c.setFillColor(INK)
    c.setFont(F_BOLD, 12.2)
    c.drawString(ML + 14, top - 21, title)
    draw_rule(c, ML + 14, top - 28, ML + CW - 14, 0.6)

    rects = [(ML + 14 + i * (inner_w + GUT), bot + 10, inner_w, top - 35 - (bot + 10))
             for i in range(ncols)]
    flow(c, fl, rects)
    return bot


def fill_gaps(c, spec, y_top, y_bot, start_idx=0, limit=5):
    """남는 자리를 박스 기사로 차례차례 메웁니다."""
    y = y_top
    for k in range(limit):
        if y - y_bot < 84:
            break
        y = filler_box(c, spec, y - (6 if k else 0), y_bot, start_idx + k) 
    return y


def all_headlines(spec, limit=60, bsize=7.8, skip=()):
    """수집된 모든 헤드라인 — 지면을 끝까지 채우는 마지막 재료."""
    rows = []
    for n in spec["names"]:
        if n in skip:
            continue
        for it in spec["results"][n]["items"]:
            w = it["when"].strftime("%m/%d") if it.get("when") else ""
            rows.append(f'[{n}] {_clean(it["title"])}  ({_clean(it["source"])} {w})')
    if not rows:
        return []
    return brief_flowables("이번 주 수집 헤드라인 전체", rows[:limit], bsize=bsize)


def page1(c, spec):
    imgs = spec.get("images", {})
    y = masthead(c, spec)
    top = spec["top"]

    y = draw_kicker(c, top.get("kicker", "이번 주 톱"), ML, y)
    y = draw_headline(c, top["headline"], ML, y - 6, CW,
                      max_size=44, min_size=21, max_lines=2)
    if top.get("sub"):
        c.setFont(F_REG, 12.6)
        c.setFillColor(GREY)
        for ln in wrap_pdf(top["sub"], F_REG, 12.6, CW)[:2]:
            y -= 17
            c.drawString(ML, y, ln)
    y -= 10
    draw_rule(c, ML, y, PW - MR, 1.0)
    y -= 12

    # 사진(왼쪽 2단) + 톱기사 본문(오른쪽 2단)
    ph_h = 215
    cap_y = place_image(c, imgs.get("top"), ML, y - ph_h, spanw(2), ph_h,
                        photo_caption(spec, "top", top.get("photo_caption", "")))
    fl = art_flowables(lead=top.get("lead"), body=top.get("body"))
    leftover, _lb = flow_balanced(c, fl, y, cap_y - 4, ncols=2, first=2)
    draw_vrule(c, colx(2) - GUT / 2, cap_y - 4, y, 0.35, rlcolors.HexColor("#b9b9bd"))

    y2 = cap_y - 15
    draw_rule(c, ML, y2, PW - MR, 0.6)
    y2 -= 12

    band = MB + 124
    sec_fl = list(leftover)
    for a in (spec["second"] or [])[:4]:
        sec_fl += art_flowables(headline=a.get("headline"), body=a.get("body"),
                                source=a.get("source"), hsize=14)
    sec_fl += _sec_briefs(spec, spec["names"], limit=7)
    sec_fl += all_headlines(spec, 40)
    _, used = flow_balanced(c, sec_fl, y2, band + 12)
    if used - (band + 12) > 86:
        fill_gaps(c, spec, used - 8, band + 12, 0)

    # ── 아래 띠 — 이번 주 한눈에
    draw_rule(c, ML, band, PW - MR, 1.2)
    c.setFillColor(INK)
    c.setFont(F_BOLD, 12)
    c.drawString(ML, band - 17, "이번 주 한눈에")
    kw = spec["keywords"][:10]
    if kw:
        c.setFont(F_REG, 9)
        c.setFillColor(GREY)
        c.drawString(ML + 80, band - 17,
                     "키워드  " + "  ·  ".join(f"{w}({n})" for w, n in kw))
    items = []
    for n in spec["names"]:
        its = spec["results"][n]["items"]
        if its:
            items.append(f'[{n}] {_clean(its[0]["title"])}')
    items = items[:12]
    per = math.ceil(len(items) / 4) or 1
    for i in range(4):
        chunk = items[i * per:(i + 1) * per]
        if not chunk:
            continue
        fl = [Paragraph(f'· {esc(t)}', _style("bb", 7.9, 11.6, F_REG, TA_LEFT,
                                              space_after=2)) for t in chunk]
        flow(c, fl, [(colx(i), MB + 20, COLW, band - 28 - (MB + 20))])
    page_foot(c, spec, 1)


def page2(c, spec):
    y = running_head(c, spec, 2, PAGE_NAMES[1])
    y = draw_headline(c, "이번 주, 세상은 이렇게 움직였다", ML, y, CW,
                      max_size=30, min_size=18, max_lines=1)
    c.setFont(F_REG, 9.4)
    c.setFillColor(GREY)
    y -= 15
    c.drawString(ML, y, "수집된 실제 기사만을 근거로 한 주간의 흐름을 정리했습니다.")
    y -= 9
    draw_rule(c, ML, y, PW - MR, 1.0)
    y -= 13

    fl = []
    for i, f3 in enumerate(spec["flow3"][:4], 1):
        fl += art_flowables(headline=f'{i}. {f3.get("title","")}',
                            body=f3.get("body", []), hsize=14)
    if spec.get("ai_text"):
        for blk in re.split(r"\n\s*\n", spec["ai_text"]):
            raw = blk.strip()
            if not raw:
                continue
            head = raw.splitlines()[0].strip()
            rest = _clean(" ".join(raw.splitlines()[1:]))
            if head.startswith(("🗞️", "📋", "🙏", "⚠️", "🔎", "▪")):
                fl.append(Paragraph(f"<b>{esc(head)}</b>",
                                    _style("hh", 11.4, 15.4, F_BOLD, TA_LEFT,
                                           space_after=4)))
                if rest:
                    fl.append(Paragraph(esc(rest),
                                        _style("bx", 8.8, 13.8, F_REG, first_indent=8.8)))
            else:
                fl.append(Paragraph(esc(_clean(raw)),
                                    _style("bx", 8.8, 13.8, F_REG, first_indent=8.8)))
    fl += brief_flowables("이번 주 키워드 풀이",
                          [f"{w} — 관련 보도 {n}건" for w, n in spec["keywords"][:12]])
    fl += _sec_briefs(spec, spec["names"], limit=4, bsize=7.9)

    tbl_top = MB + 224
    _, used = flow_balanced(c, fl, y, tbl_top + 14)
    if used - (tbl_top + 14) > 86:
        fill_gaps(c, spec, used - 8, tbl_top + 14, 1)

    draw_rule(c, ML, tbl_top, PW - MR, 0.8)
    c.setFillColor(INK)
    c.setFont(F_BOLD, 12)
    c.drawString(ML, tbl_top - 16, spec["table"].get("title", "표"))
    t = make_table(spec["table"]["headers"], spec["table"]["rows"], spanw(2))
    tw, th = t.wrapOn(c, spanw(2), tbl_top - MB - 40)
    t.drawOn(c, ML, max(MB + 22, tbl_top - 24 - th))

    c.setFillColor(INK)
    c.setFont(F_BOLD, 12)
    c.drawString(colx(2), tbl_top - 16, "섹션별 수집 현황")
    rows = [[n, str(len(spec["results"][n]["items"])), spec["results"][n]["group"]]
            for n in spec["names"]]
    rows.sort(key=lambda r: -int(r[1]))
    t2 = make_table(["섹션", "기사", "분류"], rows[:12], spanw(2))
    tw2, th2 = t2.wrapOn(c, spanw(2), tbl_top - MB - 40)
    t2.drawOn(c, colx(2), max(MB + 22, tbl_top - 24 - th2))
    page_foot(c, spec, 2)


def section_page(c, spec, no, img_key):
    names = [n for n in spec["page_sections"].get(no, [])
             if spec["results"].get(n, {}).get("items")]
    label = " · ".join(names[:3]) if names else PAGE_NAMES[no - 1]
    y = running_head(c, spec, no, label[:26])
    imgs = spec.get("images", {})

    if not names:
        c.setFillColor(GREY)
        c.setFont(F_REG, 10)
        c.drawString(ML, y - 20, "이 면에 배치된 섹션이 없습니다. "
                                 "‘면 배치’에서 섹션을 골라 주세요.")
        page_foot(c, spec, no)
        return

    first = names[0]
    items = spec["results"][first]["items"]
    head = _clean(items[0]["title"]) if items else first
    x = ML + draw_label(c, first, ML, y - 12) + 8
    c.setFillColor(GREY)
    c.setFont(F_REG, 8.4)
    c.drawString(x, y - 11, f'기사 {len(items)}건  ·  수집 {spec["period"]}')
    y -= 26
    y = draw_headline(c, head, ML, y, CW, max_size=30, min_size=17, max_lines=2)
    y -= 9
    draw_rule(c, ML, y, PW - MR, 0.8)
    y -= 12

    ph_h = 178
    cap_y = place_image(c, imgs.get(img_key), colx(2), y - ph_h, spanw(2), ph_h,
                        photo_caption(spec, img_key, f"{first} 관련."))

    ai_secs = {s.get("name"): s for s in (spec.get("ai_sections") or [])}
    a = ai_secs.get(first) or _fallback_article(first, items) or {}
    lead_fl = art_flowables(lead=a.get("lead"), body=a.get("body", []))
    flow_balanced(c, lead_fl, y, cap_y - 4, ncols=2, first=0)
    draw_vrule(c, colx(2) - GUT / 2, cap_y - 4, y, 0.35, rlcolors.HexColor("#b9b9bd"))

    y2 = cap_y - 15
    draw_rule(c, ML, y2, PW - MR, 0.6)
    y2 -= 12

    fl = brief_flowables(f"{first} — 이번 주 보도",
                         [f'{_clean(it["title"])}  ({_clean(it["source"])})'
                          for it in items[1:9]])
    for n in names[1:]:
        its = spec["results"][n]["items"]
        if not its:
            continue
        aa = ai_secs.get(n) or _fallback_article(n, its) or {}
        fl += art_flowables(label=n,
                            headline=aa.get("headline") or _clean(its[0]["title"]),
                            lead=aa.get("lead"), body=aa.get("body", []),
                            source=_clean(its[0]["source"]), hsize=13)
        fl += brief_flowables("함께 보기",
                              [f'{_clean(it["title"])}  ({_clean(it["source"])})'
                               for it in its[1:7]], bsize=7.9)
    # 지면이 남으면 다른 섹션 헤드라인으로 채운다
    others = [n for n in spec["names"] if n not in names]
    fl += _sec_briefs(spec, others[:6], limit=4, bsize=7.8)
    fl += all_headlines(spec, 60, skip=tuple(names))
    _, used = flow_balanced(c, fl, y2, MB + 26)
    if used - (MB + 26) > 86:
        fill_gaps(c, spec, used - 8, MB + 26, no)
    page_foot(c, spec, no)


def page5(c, spec):
    y = running_head(c, spec, 5, PAGE_NAMES[4])
    imgs = spec.get("images", {})
    ct = spec["cartoon"]

    y = draw_headline(c, "만평", ML, y, CW, max_size=26, min_size=18, max_lines=1)
    y -= 7
    draw_rule(c, ML, y, PW - MR, 1.0)
    y -= 12

    cw_ = spanw(3)
    _ci = imgs.get("cartoon")
    ch_ = cw_ * (_ci.size[1] / _ci.size[0]) if _ci is not None else cw_ * 0.71
    ch_ = min(ch_, y - MB - 260)
    cw_ = min(cw_, ch_ * (_ci.size[0] / _ci.size[1])) if _ci is not None else cw_
    place_image(c, _ci, ML, y - ch_, cw_, ch_, None)
    cap_y = y - ch_

    ed = spec["editorial"]
    fl = [Paragraph('<font color="#b91c1c"><b>사설</b></font>',
                    _style("k", 9.6, 13, F_BOLD, TA_LEFT)),
          Paragraph(esc(ed.get("title", "사설")),
                    _style("eh", 15, 19.5, F_BOLD, TA_LEFT, space_after=6))]
    for p_ in ed.get("body", []):
        fl.append(Paragraph(esc(p_), _style("eb", 9, 14, F_REG, first_indent=9)))
    left = flow(c, fl, [(colx(3), cap_y, COLW, y - cap_y)])
    draw_vrule(c, colx(3) - GUT / 2, cap_y, y, 0.35, rlcolors.HexColor("#b9b9bd"))

    y2 = cap_y - 14
    draw_rule(c, ML, y2, PW - MR, 0.6)
    y2 -= 12

    more = list(left)
    more += brief_flowables("만평을 이렇게 읽어 주세요",
                            [x for x in [ct.get("left_line", ""),
                                         ct.get("right_line", ""),
                                         ct.get("caption", "")] if x])
    more += brief_flowables("이번 주 묵상 질문", [
        "이 한 주간 나를 가장 흔든 소식은 무엇이었습니까?",
        "그 소식 앞에서 성경은 무엇이라 말합니까?",
        "우리 교회가 실제로 할 수 있는 한 가지는 무엇입니까?"])
    more += brief_flowables("이번 주 기도", spec["prayer"][:6])
    more += _sec_briefs(spec, spec["names"], limit=4, bsize=7.8)
    more += all_headlines(spec, 60)
    _, used = flow_balanced(c, more, y2, MB + 26)
    if used - (MB + 26) > 86:
        fill_gaps(c, spec, used - 8, MB + 26, 2)
    page_foot(c, spec, 5)


def page6(c, spec):
    y = running_head(c, spec, 6, PAGE_NAMES[5])
    imgs = spec.get("images", {})
    y = draw_headline(c, "숫자와 그림으로 보는 한 주", ML, y, CW,
                      max_size=28, min_size=18, max_lines=1)
    y -= 7
    draw_rule(c, ML, y, PW - MR, 1.0)
    y -= 12

    if imgs.get("chart") is not None:
        ci = imgs["chart"]
        ch_ = min(290, CW * ci.size[1] / ci.size[0])
        y = place_image(c, ci, ML, y - ch_, CW, ch_,
                        "섹션별 기사 수 · 이번 주 키워드 (실제 수집 자료)") - 14

    dg = imgs.get("diagram")
    bottom_zone = MB + 200
    if dg is not None:
        dh = min(spanw(2) * dg.size[1] / dg.size[0], y - bottom_zone - 24)
        cap_y = place_image(c, dg, ML, y - dh, spanw(2), dh,
                            spec["diagram"].get("title", "도해"))
    else:
        cap_y = y

    t = make_table(spec["table"]["headers"], spec["table"]["rows"], spanw(2))
    tw, th = t.wrapOn(c, spanw(2), y - bottom_zone)
    c.setFillColor(INK)
    c.setFont(F_BOLD, 11.4)
    c.drawString(colx(2), y - 12, spec["table"].get("title", "표"))
    t.drawOn(c, colx(2), max(bottom_zone, y - 20 - th))

    yb = min(cap_y, y - 20 - th) - 16
    if yb > MB + 60:
        draw_rule(c, ML, yb, PW - MR, 0.6)
        yb -= 10
        fl = brief_flowables("읽는 법",
                             ["막대가 긴 섹션일수록 이번 주 보도가 많았던 분야입니다.",
                              "키워드는 수집된 기사 제목에 실제로 등장한 낱말만 셉니다.",
                              "도해 가운데는 이번 주의 중심 사안, 둘레는 파생된 논점입니다."])
        fl += brief_flowables("이번 주 키워드",
                              [f"{w} — {n}건" for w, n in spec["keywords"][:12]],
                              bsize=8.0)
        fl += _sec_briefs(spec, spec["names"], limit=4, bsize=7.8)
        fl += all_headlines(spec, 60)
        _, used = flow_balanced(c, fl, yb, MB + 26)
        if used - (MB + 26) > 86:
            fill_gaps(c, spec, used - 8, MB + 26, 3)
    page_foot(c, spec, 6)


def page8(c, spec):
    y = running_head(c, spec, 8, PAGE_NAMES[7])
    imgs = spec.get("images", {})
    y = draw_headline(c, "이 한 주를 말씀으로 읽는다", ML, y, CW,
                      max_size=30, min_size=18, max_lines=1)
    y -= 7
    draw_rule(c, ML, y, PW - MR, 1.0)
    y -= 12

    ph_h = 158
    cap_y = place_image(c, imgs.get("p8"), colx(2), y - ph_h, spanw(2), ph_h,
                        photo_caption(spec, "p8", "강단은 세상의 소리를 말씀으로 번역하는 자리다."))

    rows = [[l.get("event", ""), l.get("text", ""), l.get("use", "")]
            for l in spec["sermon_links"][:7]]
    t = Table([[Paragraph(f"<b>{esc(h)}</b>",
                          _style("th", 8.2, 11.5, F_BOLD, TA_CENTER,
                                 color=rlcolors.white))
                for h in ["이번 주 사건", "연결 본문", "강단에서 쓰는 법"]]] +
              [[Paragraph(esc(x), _style("td", 8.2, 12, F_REG, TA_LEFT)) for x in r]
               for r in rows],
              colWidths=[spanw(2) * 0.36, spanw(2) * 0.2, spanw(2) * 0.44])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rlcolors.HexColor("#111827")),
        ("GRID", (0, 0), (-1, -1), 0.4, rlcolors.HexColor("#9ca3af")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [rlcolors.white, rlcolors.HexColor("#f3f4f6")]),
    ]))
    tw, th = t.wrapOn(c, spanw(2), y - MB)
    t.drawOn(c, ML, y - th)

    ybot = min(cap_y, y - th) - 16
    draw_rule(c, ML, ybot, PW - MR, 0.6)
    ybot -= 12

    colo = MB + 128
    fl = brief_flowables("이번 주 기도", spec["prayer"][:6])
    fl += brief_flowables("다룰 때 조심할 점", [
        "정치적으로 민감한 사안은 어느 한쪽을 편드는 표현을 피하십시오.",
        "사실관계가 확정되지 않은 사건은 단정적으로 말하지 마십시오.",
        "사망·재난은 애도와 절제된 표현으로 다루십시오."])
    fl += brief_flowables("이번 주 전체 헤드라인",
                          [f'[{n}] {_clean(it["title"])}'
                           for n in spec["names"]
                           for it in spec["results"][n]["items"][:3]][:36], bsize=7.7)
    _, used = flow_balanced(c, fl, ybot, colo + 12)
    if used - (colo + 12) > 86:
        fill_gaps(c, spec, used - 8, colo + 12, 4)

    # ── 판권
    draw_rule(c, ML, colo, PW - MR, 1.2)
    c.setFillColor(INK)
    c.setFont(F_BOLD, 12)
    c.drawString(ML, colo - 17, spec["title"])
    c.setFont(F_REG, 8.6)
    c.setFillColor(GREY)
    left_lines = [f'발행 · {spec["publisher"]}',
                  f'{spec["issue"]}   {spec["date_line"]}',
                  f'수집 기간 {spec["period"]}']
    if spec.get("editor"):
        left_lines.insert(1, f'편집 · {spec["editor"]}')
    right_lines = [f'수집 기사 {spec["total"]}건 / 섹션 {spec["nsec"]}개',
                   '기사 출처 · Google 뉴스 (각 매체 원문 링크 보유)',
                   '이 신문은 설교·목회 자료용 내부 간행물입니다.']
    yy = colo - 34
    for s in left_lines:
        c.drawString(ML, yy, s)
        yy -= 13
    yy = colo - 34
    for s in right_lines:
        c.drawString(colx(2), yy, s)
        yy -= 13
    c.setFont(F_REG, 7.6)
    c.setFillColor(SOFT)
    c.drawCentredString(PW / 2, MB + 26,
                        "MY 설교 AI 스튜디오 Pro — 시사주간뉴스 종이신문 발행 기능으로 만들었습니다.")
    page_foot(c, spec, 8)

def make_pdf(spec):
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A3)
    c.setTitle(f'{spec["title"]} {spec["date_line"]}')
    c.setAuthor(spec["publisher"])
    c.setSubject("시사주간뉴스 종이신문 (A3 세로 8면)")

    for no in range(1, 9):
        if no == 1:
            page1(c, spec)
        elif no == 2:
            page2(c, spec)
        elif no == 5:
            page5(c, spec)
        elif no == 6:
            page6(c, spec)
        elif no == 8:
            page8(c, spec)
        else:
            section_page(c, spec, no, {3: "p3", 4: "p4", 7: "p7"}[no])
        c.showPage()
    c.save()
    return buf.getvalue()


# ==============================================================================
# 5. WORD (A3 세로 · 2단)
# ==============================================================================
def make_docx(spec):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    imgs = spec.get("images", {})

    def _cols(section, n, space=360):
        sectPr = section._sectPr
        cols = sectPr.xpath("./w:cols")
        if cols:
            cc = cols[0]
        else:
            cc = OxmlElement("w:cols")
            sectPr.append(cc)
        cc.set(qn("w:num"), str(n))
        cc.set(qn("w:space"), str(space))
        cc.set(qn("w:equalWidth"), "1")

    def _setup(section):
        section.page_width = Mm(297)
        section.page_height = Mm(420)
        section.left_margin = Mm(14)
        section.right_margin = Mm(14)
        section.top_margin = Mm(13)
        section.bottom_margin = Mm(14)

    s0 = doc.sections[0]
    _setup(s0)
    _cols(s0, 1)

    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(9.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    def p(text="", size=9.5, bold=False, align=None, color=None, space_after=4,
          indent=False, italic=False):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.space_before = Pt(0)
        if indent:
            par.paragraph_format.first_line_indent = Pt(size)
        if align is not None:
            par.alignment = align
        r = par.add_run(str(text))
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return par

    def rule(weight=8, color="111111"):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        pPr = par._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(weight))
        b.set(qn("w:space"), "1")
        b.set(qn("w:color"), color)
        pbdr.append(b)
        pPr.append(pbdr)

    def img(key, width_mm, caption=None):
        im = imgs.get(key)
        if im is None:
            return
        bio = io.BytesIO(pil_to_bytes(im, "JPEG" if key != "diagram" else "PNG"))
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(bio, width=Mm(width_mm))
        if caption:
            p(caption, 8, False, WD_ALIGN_PARAGRAPH.CENTER, "666666", 8)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(h))
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row[:len(headers)]):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(8.5)
                r.font.name = "맑은 고딕"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        p("", 6)

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def head_bar(no, name):
        rule(18)
        p(f'{no}면   {name}          {spec["title"]} · {spec["date_line"]}',
          8.5, True, None, "333333", 4)
        rule(8)

    # ── 1면 제호
    p(f'{spec["date_line"]}                    {spec["publisher"]} · {spec["issue"]}',
      8.5, False, WD_ALIGN_PARAGRAPH.CENTER, "444444", 2)
    rule(24)
    p(spec["title"], 46, True, WD_ALIGN_PARAGRAPH.CENTER, "111111", 2)
    p(spec["sub"], 11, False, WD_ALIGN_PARAGRAPH.CENTER, "555555", 2)
    rule(24)
    p(f'수집 기간 {spec["period"]}  ·  기사 {spec["total"]}건  ·  섹션 {spec["nsec"]}개  '
      f'·  출처 Google 뉴스  ·  A3 세로 8면',
      8.2, False, WD_ALIGN_PARAGRAPH.CENTER, "666666", 8)

    top = spec["top"]
    p(top.get("kicker", "이번 주 톱"), 10, True, None, "B91C1C", 2)
    p(top["headline"], 26, True, None, "111111", 6)
    rule(12)
    img("top", 150, photo_caption(spec, "top", top.get("photo_caption", "")))

    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    _setup(s1)
    _cols(s1, 2, 400)

    if top.get("lead"):
        p(top["lead"], 10, True, None, None, 5)
    for b in top.get("body", []):
        p(b, 9.5, False, None, None, 4, indent=True)
    for a in (spec["second"] or [])[:3]:
        p(a.get("headline", ""), 13, True, None, "111111", 4)
        for b in a.get("body", []):
            p(b, 9.5, False, None, None, 4, indent=True)
        if a.get("source"):
            p(a["source"], 8, False, None, "6B7280", 8)

    p("이번 주 한눈에", 12, True, None, "111111", 4)
    for n in spec["names"]:
        its = spec["results"][n]["items"]
        if its:
            p(f'[{n}] {_clean(its[0]["title"])}', 8.5, False, None, None, 2)

    # ── 2면
    page_break()
    head_bar(2, PAGE_NAMES[1])
    p("이번 주, 세상은 이렇게 움직였다", 20, True, None, "111111", 6)
    for i, f3 in enumerate(spec["flow3"][:4], 1):
        p(f'{i}. {f3.get("title","")}', 13, True, None, "111111", 4)
        for b in f3.get("body", []):
            p(b, 9.5, False, None, None, 4, indent=True)
    if spec.get("ai_text"):
        p("", 6)
        for blk in re.split(r"\n\s*\n", spec["ai_text"]):
            b = _clean(blk)
            if not b:
                continue
            if b.startswith(("🗞️", "📋", "🙏", "⚠️", "🔎")):
                p(b.splitlines()[0] if "\n" in blk else b, 11.5, True, None, "111111", 4)
                rest = _clean(" ".join(blk.splitlines()[1:]))
                if rest:
                    p(rest, 9.2, False, None, None, 4, indent=True)
            else:
                p(b, 9.2, False, None, None, 4, indent=True)
    p("", 6)
    p(spec["table"].get("title", "표"), 12, True, None, "111111", 4)
    table(spec["table"]["headers"], spec["table"]["rows"])
    p("이번 주 키워드 풀이", 12, True, None, "111111", 4)
    for w, n in spec["keywords"][:10]:
        p(f'{w} — 관련 보도 {n}건', 8.8, False, None, None, 2)

    # ── 3·4·7면
    for no, ik in ((3, "p3"), (4, "p4"), (7, "p7")):
        page_break()
        names = spec["page_sections"].get(no, [])
        head_bar(no, " · ".join(names[:3]) if names else PAGE_NAMES[no - 1])
        if not names:
            p("이 면에 배치된 섹션이 없습니다.", 10, False, None, "666666")
            continue
        first = names[0]
        items = spec["results"][first]["items"]
        p(f'[{first}]  기사 {len(items)}건', 9.5, True, None, "B91C1C", 3)
        p(_clean(items[0]["title"]) if items else first, 19, True, None, "111111", 5)
        rule(10)
        img(ik, 140, photo_caption(spec, ik, f"{first} 관련."))
        ai_secs = {s.get("name"): s for s in (spec.get("ai_sections") or [])}
        a = ai_secs.get(first) or _fallback_article(first, items) or {}
        if a.get("lead"):
            p(a["lead"], 10, True, None, None, 4)
        for b in a.get("body", []):
            p(b, 9.5, False, None, None, 4, indent=True)
        p(f'{first} — 이번 주 보도', 11.5, True, None, "111111", 4)
        for i, it in enumerate(items[1:8], 1):
            p(f'{i}. {_clean(it["title"])}  ({_clean(it["source"])})', 8.6, False,
              None, None, 2)
        for n in names[1:]:
            its = spec["results"][n]["items"]
            if not its:
                continue
            p("", 4)
            p(f'[{n}]', 9.5, True, None, "B91C1C", 2)
            aa = ai_secs.get(n) or _fallback_article(n, its) or {}
            p(aa.get("headline") or _clean(its[0]["title"]), 13, True, None, "111111", 3)
            if aa.get("lead"):
                p(aa["lead"], 9.6, True, None, None, 3)
            for b in aa.get("body", []):
                p(b, 9.3, False, None, None, 3, indent=True)
            for i, it in enumerate(its[1:6], 1):
                p(f'· {_clean(it["title"])}  ({_clean(it["source"])})', 8.4, False,
                  None, None, 2)

    # ── 5면
    page_break()
    head_bar(5, PAGE_NAMES[4])
    p("만평", 20, True, None, "111111", 5)
    img("cartoon", 150, spec["cartoon"].get("caption"))
    ed = spec["editorial"]
    p("사설", 10, True, None, "B91C1C", 2)
    p(ed.get("title", "사설"), 15, True, None, "111111", 5)
    for b in ed.get("body", []):
        p(b, 9.5, False, None, None, 4, indent=True)
    p("이번 주 묵상 질문", 12, True, None, "111111", 4)
    for q in ["이 한 주간 나를 가장 흔든 소식은 무엇이었습니까?",
              "그 소식 앞에서 성경은 무엇이라 말합니까?",
              "우리 교회가 실제로 할 수 있는 한 가지는 무엇입니까?"]:
        p(f'· {q}', 9.2, False, None, None, 3)

    # ── 6면
    page_break()
    head_bar(6, PAGE_NAMES[5])
    p("숫자와 그림으로 보는 한 주", 20, True, None, "111111", 6)
    if imgs.get("chart") is not None:
        img("chart", 150, "섹션별 기사 수 · 이번 주 키워드 (실제 수집 자료)")
    img("diagram", 140, spec["diagram"].get("title", "도해"))
    p(spec["table"].get("title", "표"), 12, True, None, "111111", 4)
    table(spec["table"]["headers"], spec["table"]["rows"])
    p("읽는 법", 11.5, True, None, "111111", 4)
    for s in ["막대가 긴 섹션일수록 이번 주 보도가 많았던 분야입니다.",
              "키워드는 수집된 기사 제목에 실제로 등장한 낱말만 셉니다.",
              "도해 가운데는 이번 주의 중심 사안, 둘레는 파생된 논점입니다."]:
        p(f'· {s}', 9, False, None, None, 3)

    # ── 8면
    page_break()
    head_bar(8, PAGE_NAMES[7])
    p("이 한 주를 말씀으로 읽는다", 20, True, None, "111111", 6)
    img("p8", 140, photo_caption(spec, "p8", "강단은 세상의 소리를 말씀으로 번역하는 자리다."))
    table(["이번 주 사건", "연결 본문", "강단에서 쓰는 법"],
          [[l.get("event", ""), l.get("text", ""), l.get("use", "")]
           for l in spec["sermon_links"][:6]])
    p("이번 주 기도", 12, True, None, "111111", 4)
    for s in spec["prayer"][:6]:
        p(f'· {s}', 9.2, False, None, None, 3)
    p("다룰 때 조심할 점", 12, True, None, "111111", 4)
    for s in ["정치적으로 민감한 사안은 어느 한쪽을 편드는 표현을 피하십시오.",
              "사실관계가 확정되지 않은 사건은 단정적으로 말하지 마십시오.",
              "사망·재난은 애도와 절제된 표현으로 다루십시오."]:
        p(f'· {s}', 9.2, False, None, None, 3)
    p("이번 주 전체 헤드라인", 12, True, None, "111111", 4)
    for n in spec["names"]:
        for it in spec["results"][n]["items"][:3]:
            p(f'[{n}] {_clean(it["title"])}  ({_clean(it["source"])})', 8.2,
              False, None, None, 2)
    rule(18)
    p(spec["title"], 13, True, None, "111111", 2)
    lines = [f'발행 · {spec["publisher"]}',
             f'{spec["issue"]}   {spec["date_line"]}',
             f'수집 기간 {spec["period"]}',
             f'수집 기사 {spec["total"]}건 / 섹션 {spec["nsec"]}개',
             '기사 출처 · Google 뉴스 (각 매체 원문 링크 보유)',
             '이 신문은 설교·목회 자료용 내부 간행물입니다.']
    if spec.get("editor"):
        lines.insert(1, f'편집 · {spec["editor"]}')
    for s in lines:
        p(s, 8.4, False, None, "666666", 2)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ==============================================================================
# 6. 면 구성 미리보기(글)
# ==============================================================================
def outline_text(spec):
    ps = spec["page_sections"]
    rows = [
        (1, "종합", f'톱기사 · 사진 · 주요기사 {len(spec["second"])}건 · 이번 주 한눈에'),
        (2, "이슈 해설", f'이번 주 큰 흐름 {len(spec["flow3"])}가지 · 표 · 키워드'),
        (3, " · ".join(ps.get(3, [])) or "—", "머리기사 · 사진 · 섹션 기사"),
        (4, " · ".join(ps.get(4, [])) or "—", "머리기사 · 사진 · 섹션 기사"),
        (5, "만평 · 사설", "한 컷 만평 · 사설 · 묵상 질문 · 기도"),
        (6, "도해 · 통계", "통계 그래프 · 도해 · 표"),
        (7, " · ".join(ps.get(7, [])) or "—", "머리기사 · 사진 · 섹션 기사"),
        (8, "신앙 · 목회", "설교 연결 포인트 표 · 기도 · 전체 헤드라인 · 판권"),
    ]
    return rows
